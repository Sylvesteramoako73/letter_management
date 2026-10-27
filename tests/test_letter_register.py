import io
import zipfile

import pymupdf
from docx import Document

from test_workflow import client, head_approves, login, respond, route_as_md  # noqa: F401  (pytest fixture)


def register(client, department_id, sender="Ministry of Finance", **extra):
    """The front desk registers a letter and the MD sends it to the department; leaves the front desk signed in."""
    login(client, "frontdesk")
    response = client.post(
        "/letters/upload",
        data={
            "sender": sender,
            "subject": "Tax clearance request",
            "received_date": "2026-10-10",
            "department_id": str(department_id),
            "letter_date": "2026-10-08",
            "agency_reference": "MOF/TAX/22",
            "signatory": "K. Mensah",
            "cc_parties": "Ghana Revenue Authority\nAuditor-General",
            "received_via": "hand_delivery",
            "action_required": "reply",
            "document": (io.BytesIO(b"%PDF-1.4"), "scan.pdf"),
            **extra,
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 201, response.get_json()
    letter = route_as_md(client, response.get_json()["id"], department_id)
    login(client, "frontdesk")
    return letter


def approve_as_md(client, letter_id):
    login(client, "md")
    response = client.post(f"/letters/{letter_id}/md-approve", json={})
    assert response.status_code == 200, response.get_json()
    return response.get_json()


def test_register_fields_are_locked_except_for_admins(client):
    client, department_id, _ = client
    letter = register(client, department_id)
    assert letter["letter_date"] == "2026-10-08"
    assert letter["signatory"] == "K. Mensah"
    assert letter["cc_parties"] == "Ghana Revenue Authority\nAuditor-General"
    assert (letter["received_via"], letter["action_required"]) == ("hand_delivery", "reply")
    assert letter["matter_status"] == "open"
    login(client, "frontdesk")
    bad = client.post(
        "/letters/upload",
        data={"sender": "A", "subject": "B", "signatory": "K. Mensah", "received_date": "2026-10-10", "department_id": str(department_id),
              "received_via": "pigeon", "document": (io.BytesIO(b"%PDF-1.4"), "scan.pdf")},
        content_type="multipart/form-data",
    )
    assert bad.status_code == 400

    for username in ("frontdesk", "head", "member"):
        login(client, username)
        assert client.patch(f"/letters/{letter['id']}/details", json={"signatory": "X"}).status_code == 403

    login(client, "mdpa")
    response = client.patch(f"/letters/{letter['id']}/details", json={"signatory": "Kofi Mensah", "received_via": "email"})
    assert response.status_code == 200
    assert (response.get_json()["signatory"], response.get_json()["received_via"]) == ("Kofi Mensah", "email")
    audit = client.get(f"/letters/{letter['id']}").get_json()["audit"]
    assert audit[-1]["action"] == "register_details_edited"
    assert client.patch(f"/letters/{letter['id']}/details", json={"subject": ""}).status_code == 400


def test_respond_creates_linked_outgoing_letter_and_approval_deletes_word(client):
    client, department_id, users = client
    letter = register(client, department_id)
    login(client, "head")
    client.post(f"/letters/{letter['id']}/assign", json={"assignee_id": users["member"]})

    login(client, "member")
    response = client.post(f"/letters/{letter['id']}/respond")
    assert response.status_code == 200
    reply = response.get_json()
    assert reply["id"] != letter["id"]
    assert reply["direction"] == "outgoing" and reply["in_reply_to"] == letter["id"]
    assert reply["our_reference"] == "SIGL/MOF/AD/0001"
    assert reply["addressee"] == "K. Mensah"
    assert reply["has_word_draft"] is True
    # Pressing Respond again opens the same draft instead of starting another.
    assert client.post(f"/letters/{letter['id']}/respond").get_json()["id"] == reply["id"]
    assert client.get(f"/letters/{letter['id']}").get_json()["letter"]["status"] == "in_progress"

    draft = client.get(f"/letters/{reply['id']}/response-docx")
    assert draft.headers["Content-Disposition"].endswith("SIGL-MOF-AD-0001.docx")
    document = Document(io.BytesIO(draft.data))
    header, footer = document.sections[0].header, document.sections[0].footer
    # The company letterhead artwork is the header and footer; our reference sits under it.
    assert len(header._element.xpath(".//pic:pic")) == 1 and len(footer._element.xpath(".//pic:pic")) == 1
    assert header.paragraphs[1].text == "Our Ref: SIGL/MOF/AD/0001"
    body = "\n".join(p.text for p in document.paragraphs)
    assert "K. Mensah\nMinistry of Finance" in body and "RE: TAX CLEARANCE REQUEST" in body

    # The letter is edited in the app: it starts from the letterhead template.
    assert reply["body_html"].startswith("<p>Date:") and "[Write the body of the letter here.]" in reply["body_html"]
    # An untouched draft cannot go for approval.
    assert client.post(f"/letters/{reply['id']}/submit").status_code == 400
    assert client.put(f"/letters/{reply['id']}/response", json={"body_html": "<p><br></p>"}).status_code == 400
    saved = client.put(
        f"/letters/{reply['id']}/response",
        json={"body_html": '<p class="ql-align-justify" onclick="x()">We confirm our <strong>tax status</strong> is current.</p>'
                           '<script>alert(1)</script><ul><li>Certificate attached</li></ul>'},
    ).get_json()
    assert saved["body_html"] == ('<p class="ql-align-justify">We confirm our <strong>tax status</strong> is current.</p>'
                                  "<ul><li>Certificate attached</li></ul>")
    assert "tax status is current" in saved["response_text"]
    word = Document(io.BytesIO(client.get(f"/letters/{reply['id']}/response-docx").data))
    paragraphs = [(p.style.name, p.text) for p in word.paragraphs]
    assert paragraphs[:2] == [("Normal", "We confirm our tax status is current."), ("List Bullet", "Certificate attached")]
    # The app adds the sign-off for the signatory (the MD by default) below the text.
    assert ("Normal", "Managing Director") in paragraphs and ("Normal", "For: Swami India Ghana Limited") in paragraphs
    assert client.post(f"/letters/{reply['id']}/submit").status_code == 200
    # Once submitted, the writer can no longer change it; the reviewing head can.
    assert client.put(f"/letters/{reply['id']}/response", json={"body_html": "<p>Changed</p>"}).status_code == 403
    login(client, "head")
    corrected = client.put(f"/letters/{reply['id']}/response", json={"body_html": saved["body_html"] + "<p>Yours faithfully,</p>"})
    assert corrected.status_code == 200
    head_approves(client, reply["id"])
    login(client, "mdpa")
    assert client.post(f"/letters/{reply['id']}/route-to-md").status_code == 200
    login(client, "md")
    approved = client.post(f"/letters/{reply['id']}/approve").get_json()
    assert approved["status"] == "finalized" and approved["dispatch_status"] == "pending"
    assert approved["has_word_draft"] is False
    assert client.get(f"/letters/{reply['id']}/response-docx").status_code == 404
    assert client.get(f"/letters/{reply['id']}/pdf").headers["Content-Disposition"].endswith("SIGL-MOF-AD-0001.pdf")

    received = client.get(f"/letters/{letter['id']}").get_json()["letter"]
    assert (received["status"], received["matter_status"]) == ("finalized", "responded")
    login(client, "frontdesk")
    assert any(n["kind"] == "dispatch" for n in client.get("/notifications").get_json())


def test_new_letter_keeps_approved_word_and_tracks_dispatch_and_delivery(client):
    client, department_id, users = client
    login(client, "member")
    created = client.post(
        "/letters/outgoing",
        json={"recipient": "Lands Commission", "subject": "Site visit", "addressee": "The Regional Lands Officer",
              "recipient_address": "P.O. Box 1\nAccra", "cc_parties": "Civil department"},
    )
    assert created.status_code == 201
    letter = created.get_json()
    assert letter["our_reference"] == "SIGL/LC/AD/0001" and letter["in_reply_to"] is None
    assert client.post(f"/letters/{letter['id']}/dispatch", json={"method": "courier"}).status_code == 403
    assert client.put(f"/letters/{letter['id']}/response", json={"response_text": "Please receive our team on Monday."}).status_code == 200

    approve_as_md(client, letter["id"])
    word = client.get(f"/letters/{letter['id']}/response-docx")
    assert word.status_code == 200
    assert "Approved by Managing Director" in "\n".join(p.text for p in Document(io.BytesIO(word.data)).paragraphs)

    # The front desk sees approved letters waiting to go out, and nothing else of the department's.
    login(client, "frontdesk")
    listed = {row["id"]: row for row in client.get("/letters").get_json()}
    assert letter["id"] in listed and listed[letter["id"]]["dispatch_status"] == "pending"
    assert client.post(f"/letters/{letter['id']}/delivery", data={"outcome": "delivered"}).status_code == 403
    assert client.post(f"/letters/{letter['id']}/dispatch", json={"method": "pigeon"}).status_code == 400
    sent = client.post(f"/letters/{letter['id']}/dispatch", json={"method": "dispatch_rider", "sent_on": "2026-10-12", "note": "Rider: Yaw"})
    assert sent.status_code == 200
    assert (sent.get_json()["dispatch_status"], sent.get_json()["dispatched_on"]) == ("sent", "2026-10-12")
    delivered = client.post(
        f"/letters/{letter['id']}/delivery",
        data={"outcome": "delivered", "delivered_on": "2026-10-13", "note": "Received by A. Owusu",
              "proof": (io.BytesIO(b"%PDF-1.4"), "stamped.pdf")},
        content_type="multipart/form-data",
    )
    assert delivered.status_code == 200
    result = delivered.get_json()
    assert (result["dispatch_status"], result["matter_status"], result["has_delivery_proof"]) == ("delivered", "open", True)
    assert client.get(f"/letters/{letter['id']}/delivery-proof").data == b"%PDF-1.4"
    assert client.post(f"/letters/{letter['id']}/dispatch", json={"method": "courier"}).status_code == 403

    login(client, "member")
    kinds = {n["kind"] for n in client.get("/notifications").get_json()}
    assert {"dispatched", "delivery"} <= kinds
    login(client, "md")
    with zipfile.ZipFile(io.BytesIO(client.get(f"/letters/{letter['id']}/folder.zip").data)) as archive:
        names = archive.namelist()
    assert {"SIGL-LC-AD-0001/SIGL-LC-AD-0001.docx", "SIGL-LC-AD-0001/SIGL-LC-AD-0001.pdf",
            "SIGL-LC-AD-0001/Proof of delivery - stamped.pdf"} <= set(names)


def test_new_letter_can_answer_a_received_letter(client):
    client, department_id, users = client
    letter = register(client, department_id)
    login(client, "member")
    assert client.post(
        "/letters/outgoing", json={"addressee": "The Director", "recipient": "Ministry of Finance", "subject": "Follow-up", "in_reply_to": letter["id"]}
    ).status_code == 403
    login(client, "head")
    client.post(f"/letters/{letter['id']}/assign", json={"assignee_id": users["member"]})
    login(client, "member")
    reply = client.post(
        "/letters/outgoing", json={"addressee": "The Director", "recipient": "Ministry of Finance", "subject": "Follow-up", "in_reply_to": letter["id"]}
    ).get_json()
    assert reply["in_reply_to"] == letter["id"]
    detail = client.get(f"/letters/{reply['id']}").get_json()
    assert detail["in_reply_to_letter"]["reference"] == letter["reference"]


def test_our_reference_serial_counts_per_institute(client):
    client, department_id, users = client
    login(client, "mdpa")
    departments = client.get("/departments").get_json()
    finance = next(d for d in departments if d["name"] == "Finance and Accounts")
    assert finance["code"] == "FA" and finance["code_is_default"]
    assert client.patch(f"/departments/{finance['id']}", json={"code": "ac"}).get_json()["code"] == "AC"
    login(client, "head")
    assert client.patch(f"/departments/{finance['id']}", json={"code": "X"}).status_code == 403

    first = register(client, department_id)
    login(client, "md")
    assert client.post(f"/letters/{first['id']}/respond").get_json()["our_reference"] == "SIGL/MOF/AD/0001"

    login(client, "financeandaccountsmember")
    created = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "Ministry of Finance", "subject": "Payment schedule"}).get_json()
    assert created["our_reference"] == "SIGL/MOF/AC/0002"
    assert created["department_id"] == finance["id"]
    assert created["assigned_to"] == users["financeandaccountsmember"]
    assert client.get(f"/letters/{created['id']}/source").status_code == 404
    other = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "Lands Commission", "subject": "Site visit"}).get_json()
    assert other["our_reference"] == "SIGL/LC/AC/0001"

    login(client, "frontdesk")
    # The front desk can register a letter to be sent, but must say which department sends it.
    assert client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "B"}).status_code == 400


def test_tracking_status_remarks_action_and_deadline(client):
    client, department_id, users = client
    letter = register(client, department_id)
    assert client.post(f"/letters/{letter['id']}/matter-status", json={"matter_status": "closed"}).status_code == 403
    login(client, "financeandaccountshead")
    assert client.post(f"/letters/{letter['id']}/matter-status", json={"matter_status": "closed"}).status_code == 403
    login(client, "head")
    assert client.post(f"/letters/{letter['id']}/matter-status", json={"matter_status": "pending"}).status_code == 400
    response = client.post(
        f"/letters/{letter['id']}/matter-status",
        json={"matter_status": "closed", "remarks": "For information only", "action_required": "information", "due_date": "2026-11-01"},
    )
    assert response.status_code == 200
    updated = response.get_json()
    assert (updated["matter_status"], updated["action_required"], updated["due_date"]) == ("closed", "information", "2026-11-01")
    listed = client.get("/letters").get_json()[0]
    assert listed["status_remarks"] == "For information only"
    assert client.get(f"/letters/{letter['id']}").get_json()["audit"][-1]["action"] == "tracking_updated"

    login(client, "member")
    outgoing = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "B"}).get_json()
    assert client.post(f"/letters/{outgoing['id']}/matter-status", json={"matter_status": "responded"}).status_code == 400


def test_received_letter_folder_includes_replies(client):
    client, department_id, _ = client
    letter = register(client, department_id)
    login(client, "md")
    reply = client.post(f"/letters/{letter['id']}/md-approve", json={"response_text": "Noted with thanks."}).get_json()
    response = client.get(f"/letters/{letter['id']}/folder.zip")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        names = archive.namelist()
        details = archive.read(f"{letter['reference']}/Register details.txt").decode()
    assert f"{letter['reference']}/Original - scan.pdf" in names
    assert f"{letter['reference']}/Responses/{reply['our_reference'].replace('/', '-')}.pdf" in names
    assert "Signed by: K. Mensah" in details
    login(client, "financeandaccountshead")
    assert client.get(f"/letters/{letter['id']}/folder.zip").status_code == 403


def test_outgoing_line_mirrors_incoming(client):
    """Register an outgoing letter, assign a writer, then department head, MD PA and MD approval."""
    client, department_id, users = client
    login(client, "frontdesk")
    created = client.post(
        "/letters/outgoing",
        json={"addressee": "The Director", "recipient": "Lands Commission", "subject": "Site visit", "department_id": department_id,
              "notes": "Ask for a visit next week", "priority": "high"},
    )
    assert created.status_code == 201
    letter = created.get_json()
    assert (letter["status"], letter["assigned_to"], letter["priority"]) == ("intake", None, "high")
    assert client.get(f"/letters/{letter['id']}").status_code == 200

    login(client, "head")
    assert any(n["kind"] == "new_letter" and n["letter_id"] == letter["id"] for n in client.get("/notifications").get_json())
    assigned = client.post(f"/letters/{letter['id']}/assign", json={"assignee_id": users["member"]}).get_json()
    assert (assigned["status"], assigned["assigned_to"]) == ("in_progress", users["member"])

    login(client, "member")
    assert client.put(f"/letters/{letter['id']}/response", json={"response_text": "Please receive our team."}).status_code == 200
    assert client.post(f"/letters/{letter['id']}/department-approve").status_code == 403
    assert client.post(f"/letters/{letter['id']}/submit").get_json()["status"] == "dept_review"

    login(client, "mdpa")
    assert client.post(f"/letters/{letter['id']}/route-to-md").status_code == 403
    login(client, "financeandaccountshead")
    assert client.post(f"/letters/{letter['id']}/department-approve").status_code == 403

    login(client, "head")
    returned = client.post(f"/letters/{letter['id']}/request-changes", json={"details": "Add the date."})
    assert returned.get_json()["status"] == "changes_requested"
    login(client, "member")
    assert any(n["kind"] == "changes" for n in client.get("/notifications").get_json())
    client.put(f"/letters/{letter['id']}/response", json={"response_text": "Please receive our team on 20 October."})
    client.post(f"/letters/{letter['id']}/submit")
    head_approves(client, letter["id"])
    login(client, "mdpa")
    assert client.post(f"/letters/{letter['id']}/route-to-md").status_code == 200
    login(client, "md")
    assert client.post(f"/letters/{letter['id']}/approve").get_json()["status"] == "finalized"


def test_department_head_can_write_and_skips_own_review(client):
    client, department_id, users = client
    login(client, "frontdesk")
    letter = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "B", "department_id": department_id}).get_json()
    login(client, "head")
    assert client.post(f"/letters/{letter['id']}/assign", json={"assignee_id": users["head"]}).get_json()["assigned_to"] == users["head"]
    client.put(f"/letters/{letter['id']}/response", json={"response_text": "Body"})
    assert client.post(f"/letters/{letter['id']}/submit").get_json()["status"] == "submitted_for_review"

    # A head can also name a team member as the writer when creating the letter.
    created = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "C", "writer": str(users["member"])}).get_json()
    assert (created["status"], created["assigned_to"]) == ("in_progress", users["member"])
    other_department = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "D", "writer": str(users["civilmember"])})
    assert other_department.status_code == 403


def test_autosaves_merge_into_one_version_per_session(client):
    client, department_id, users = client
    login(client, "member")
    letter = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "B"}).get_json()
    for text in ("First draft", "Second draft", "Third draft"):
        assert client.put(f"/letters/{letter['id']}/response", json={"body_html": f"<p>{text}</p>"}).status_code == 200
    detail = client.get(f"/letters/{letter['id']}").get_json()
    assert [v["change_summary"] for v in detail["versions"]] == ["Letter edited in the app", "Letter started on letterhead"]
    assert [e["action"] for e in detail["audit"]] == ["outgoing_created", "response_saved"]
    assert detail["letter"]["body_html"] == "<p>Third draft</p>"
    # Someone else's edit starts its own version, once the writer has closed the editor.
    client.delete(f"/letters/{letter['id']}/editing")
    login(client, "md")
    assert client.put(f"/letters/{letter['id']}/response", json={"body_html": "<p>MD edit</p>"}).status_code == 200
    assert len(client.get(f"/letters/{letter['id']}").get_json()["versions"]) == 3


def test_front_desk_writes_and_responds_like_everyone_else(client):
    client, department_id, users = client
    login(client, "frontdesk")
    own = client.post(
        "/letters/outgoing", json={"addressee": "The Director", "recipient": "Lands Commission", "subject": "Site visit", "department_id": department_id, "writer": "self"}
    ).get_json()
    assert (own["status"], own["assigned_to"]) == ("in_progress", users["frontdesk"])
    assert client.put(f"/letters/{own['id']}/response", json={"response_text": "Please receive our team."}).status_code == 200
    # The front desk's draft goes to the sending department's head, like a team member's.
    assert client.post(f"/letters/{own['id']}/submit").get_json()["status"] == "dept_review"

    # A received letter can be answered straight away, before anyone is assigned.
    letter = register(client, department_id)
    reply = client.post(f"/letters/{letter['id']}/respond")
    assert reply.status_code == 200 and reply.get_json()["assigned_to"] == users["frontdesk"]
    assert client.get(f"/letters/{letter['id']}").get_json()["letter"]["status"] == "in_progress"

    # Responding needs both: seeing the letter and being allowed to write.
    login(client, "financeandaccountsmember")
    assert client.post(f"/letters/{letter['id']}/respond").status_code == 403


def test_letters_close_only_on_the_stamped_copy(client):
    client, department_id, users = client
    letter = register(client, department_id, signatory_title="Chief Director")
    assert letter["signatory_title"] == "Chief Director"
    login(client, "head")
    client.post(f"/letters/{letter['id']}/assign", json={"assignee_id": users["member"]})
    login(client, "member")
    reply = client.post(f"/letters/{letter['id']}/respond").get_json()
    # The reply goes to whoever signed, and their reference sits under ours.
    assert (reply["addressee"], reply["addressee_title"]) == ("K. Mensah", "Chief Director")
    header = Document(io.BytesIO(client.get(f"/letters/{reply['id']}/response-docx").data)).sections[0].header
    assert [p.text for p in header.paragraphs[1:3]] == ["Our Ref: SIGL/MOF/AD/0001", "Your Ref: MOF/TAX/22"]
    assert "Your Ref" not in reply["body_html"] and "Chief Director" in reply["body_html"]
    client.put(f"/letters/{reply['id']}/response", json={"body_html": "<p>We confirm our tax status.</p>"})
    approve_as_md(client, reply["id"])
    page = pymupdf.open(stream=client.get(f"/letters/{reply['id']}/pdf").data, filetype="pdf")[0]
    corner = page.get_text(clip=pymupdf.Rect(300, 120, 560, 170))
    assert "Our Ref: SIGL/MOF/AD/0001" in corner and "Your Ref: MOF/TAX/22" in corner

    login(client, "head")
    # Neither letter can be marked closed before the stamped copy of the reply is in.
    refused = client.post(f"/letters/{letter['id']}/matter-status", json={"matter_status": "closed"})
    assert refused.status_code == 400 and "SIGL/MOF/AD/0001" in refused.get_json()["error"]
    assert client.post(f"/letters/{reply['id']}/matter-status", json={"matter_status": "closed"}).status_code == 400
    login(client, "frontdesk")
    client.post(f"/letters/{reply['id']}/dispatch", json={"method": "dispatch_rider"})
    delivered = client.post(f"/letters/{reply['id']}/delivery", data={"outcome": "delivered"}).get_json()
    assert (delivered["dispatch_status"], delivered["matter_status"]) == ("delivered", "open")

    # The stamped copy does not close anything by itself; closing is marked.
    stamped = client.post(
        f"/letters/{reply['id']}/delivery",
        data={"outcome": "delivered", "proof": (io.BytesIO(b"%PDF-1.4"), "stamped.pdf")},
        content_type="multipart/form-data",
    ).get_json()
    assert (stamped["matter_status"], stamped["has_delivery_proof"]) == ("open", True)
    login(client, "head")
    assert client.get(f"/letters/{letter['id']}").get_json()["letter"]["matter_status"] == "responded"
    for letter_id in (reply["id"], letter["id"]):
        closed = client.post(f"/letters/{letter_id}/matter-status", json={"matter_status": "closed"})
        assert closed.status_code == 200 and closed.get_json()["matter_status"] == "closed"


def test_md_opens_and_closes_any_letter(client):
    client, department_id, _ = client
    login(client, "member")
    letter = client.post("/letters/outgoing", json={"addressee": "The Director", "recipient": "A", "subject": "B"}).get_json()
    assert client.post(f"/letters/{letter['id']}/matter-status", json={"matter_status": "closed"}).status_code == 400
    login(client, "md")
    for status in ("closed", "open", "closed"):
        changed = client.post(f"/letters/{letter['id']}/matter-status", json={"matter_status": status})
        assert changed.status_code == 200 and changed.get_json()["matter_status"] == status


def test_signer_and_addressee_are_required(client):
    client, department_id, _ = client
    login(client, "frontdesk")
    missing = client.post(
        "/letters/upload",
        data={"sender": "A", "subject": "B", "received_date": "2026-10-10", "department_id": str(department_id),
              "document": (io.BytesIO(b"%PDF-1.4"), "scan.pdf")},
        content_type="multipart/form-data",
    )
    assert missing.status_code == 400 and "signed" in missing.get_json()["error"]
    login(client, "member")
    assert client.post("/letters/outgoing", json={"recipient": "A", "subject": "B"}).status_code == 400
