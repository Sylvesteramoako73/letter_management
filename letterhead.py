"""Company letterhead: reference codes, the letter template, and the formats a letter is kept in.

A letter is written in the in-app editor and stored as a small, sanitised subset of HTML.
The Word copy and the final PDF are both produced from that HTML.
"""

import re
from datetime import datetime, timezone
from html import escape
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

COMPANY = "SWAMI INDIA GHANA LIMITED"
COMPANY_CODE = "SIGL"

# The company's letterhead artwork (A4). Header and footer images are cut from it for Word and the editor.
LETTERHEAD_DIR = Path(__file__).resolve().parent / "static" / "letterhead"
LETTERHEAD_PDF = str(LETTERHEAD_DIR / "sigl-letterhead.pdf")
HEADER_PNG = LETTERHEAD_DIR / "header.png"
FOOTER_PNG = LETTERHEAD_DIR / "footer.png"
# In points on the A4 page: the header rule sits at y=129 and the footer rule at y=811.
REFERENCE_AREA = (300, 133, 554, 163)
BODY_AREA = (60, 166, 550, 795)
PLACEHOLDER = "[Write the body of the letter here.]"
LETTER_FONT = "Times New Roman"
_MINOR_WORDS = {"and", "of", "the", "&", "for"}


def institute_code(name: str) -> str:
    """First letter of every word: "Ministry of Finance" -> "MOF"."""
    words = re.findall(r"[A-Za-z0-9]+", name or "")
    return "".join(word[0] for word in words).upper()[:8] or "GEN"


def department_code(name: str) -> str:
    """Default short name when an admin has not set one: "Finance and Accounts" -> "FA", "Civil" -> "CI"."""
    words = [word for word in re.findall(r"[A-Za-z0-9]+", name or "") if word.lower() not in _MINOR_WORDS]
    if not words:
        return "GEN"
    if len(words) == 1:
        return words[0][:2].upper()
    return "".join(word[0] for word in words).upper()[:4]


# ---------------------------------------------------------------- The letter as HTML


def template_html(letter, body_text: str | None = None) -> str:
    """The text the writer edits: date, address block, subject, body and "Yours faithfully".

    The sign-off below it (signature, name, title) is added from the letter's signatory by
    closing_html, so changing who signs never means retyping the end of the letter.
    """
    reply = bool(letter["in_reply_to"])
    lines = [f"<p>Date: {datetime.now(timezone.utc).strftime('%d %B %Y')}</p>", "<p><br></p>"]
    # Their reference sits in the corner under ours, so it is not repeated here.
    address = [letter["addressee"], letter["addressee_title"], letter["sender"], *(letter["recipient_address"] or "").splitlines()]
    lines += [f"<p>{escape(part.strip())}</p>" for part in address if part and part.strip()]
    subject = f"RE: {letter['subject']}" if reply else letter["subject"]
    lines += ["<p><br></p>", "<p>Dear Sir/Madam,</p>", f"<p><strong><u>{escape(subject.upper())}</u></strong></p>"]
    blocks = [block.strip() for block in (body_text or "").replace("\r\n", "\n").split("\n\n") if block.strip()]
    lines += [f'<p class="ql-align-justify">{escape(block).replace(chr(10), "<br>")}</p>' for block in blocks or [PLACEHOLDER]]
    lines += ["<p><br></p>", "<p>Yours faithfully,</p>"]
    return "".join(lines)


def closing_html(signer_name: str, signer_title: str, signature_html: str = "", enclosure_names=(), cc_parties: str | None = None) -> str:
    """Signature, signatory, then any enclosures and copies, in the usual order for a letter."""
    lines = [signature_html or "<p><br></p><p><br></p>"]
    # The MD's account may just be called "Managing Director"; do not print the title twice.
    lines.append(f"<p><strong>{escape(signer_name)}</strong></p>")
    if signer_title and signer_title != signer_name:
        lines.append(f"<p>{escape(signer_title)}</p>")
    lines.append(f"<p>For: {COMPANY.title()}</p>")
    if enclosure_names:
        items = "".join(f"<li>{escape(name)}</li>" for name in enclosure_names)
        lines.append(f"<p><br></p><p><strong>Enc.:</strong></p><ol>{items}</ol>")
    parties = [line.strip() for line in (cc_parties or "").splitlines() if line.strip()]
    if parties:
        lines.append("<p><br></p><p>cc:</p>" + "".join(f"<p>{escape(party)}</p>" for party in parties))
    return "".join(lines)


_BLOCKS = {"p", "h1", "h2", "h3", "li", "blockquote"}
_INLINE = {"strong", "b", "em", "i", "u", "s", "span"}
_ALLOWED = _BLOCKS | _INLINE | {"ul", "ol", "br"}
_CLASS = re.compile(r"^ql-(align-(center|right|justify)|indent-[1-8]|size-(small|large|huge))$")
_DROP_CONTENT = {"script", "style", "head", "title", "iframe", "object"}


class _Sanitiser(HTMLParser):
    """Keeps formatting tags and Quill's alignment, indent and size classes; drops everything else."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out = []
        self.open = []
        self.skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in _DROP_CONTENT:
            self.skipping += 1
            return
        if self.skipping or tag not in _ALLOWED:
            return
        if tag == "br":
            self.out.append("<br>")
            return
        classes = " ".join(c for c in (dict(attrs).get("class") or "").split() if _CLASS.match(c))
        self.out.append(f'<{tag} class="{classes}">' if classes else f"<{tag}>")
        self.open.append(tag)

    def handle_startendtag(self, tag, attrs):
        if tag == "br" and not self.skipping:
            self.out.append("<br>")

    def handle_endtag(self, tag):
        if tag in _DROP_CONTENT:
            self.skipping = max(0, self.skipping - 1)
            return
        if self.skipping or tag not in self.open:
            return
        # Close anything left open inside it so the output stays well formed.
        while self.open:
            current = self.open.pop()
            # The editor saves a blank line as an empty paragraph, which would collapse to nothing in the PDF.
            if current in _BLOCKS and self.out and self.out[-1].startswith(f"<{current}"):
                self.out.append("<br>")
            self.out.append(f"</{current}>")
            if current == tag:
                break

    def handle_data(self, data):
        if not self.skipping:
            self.out.append(escape(data.replace("\xa0", " "), quote=False))

    def result(self):
        while self.open:
            self.out.append(f"</{self.open.pop()}>")
        return "".join(self.out)


def sanitize_html(html: str) -> str:
    parser = _Sanitiser()
    parser.feed(html or "")
    parser.close()
    return parser.result()


class _Blocks(HTMLParser):
    """Reads sanitised HTML into paragraphs of formatted runs, for Word and plain text."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks = []
        self.lists = []
        self.styles = []
        self.current = None

    def _block(self, tag, attrs):
        classes = (dict(attrs).get("class") or "").split()
        kind = self.lists[-1] if tag == "li" and self.lists else tag
        self.current = {"tag": tag, "list": kind if tag == "li" else None, "classes": classes, "runs": []}
        self.blocks.append(self.current)

    def handle_starttag(self, tag, attrs):
        if tag in {"ul", "ol"}:
            self.lists.append(tag)
        elif tag in _BLOCKS:
            self._block(tag, attrs)
        elif tag == "br":
            if self.current is None:
                self._block("p", [])
            self.current["runs"].append(("\n", set()))
        elif tag in _INLINE:
            classes = (dict(attrs).get("class") or "").split()
            self.styles.append({"strong": "b", "b": "b", "em": "i", "i": "i", "u": "u", "s": "s"}.get(tag) or next(
                (c for c in classes if c.startswith("ql-size-")), ""))

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)

    def handle_endtag(self, tag):
        if tag in {"ul", "ol"} and self.lists:
            self.lists.pop()
        elif tag in _BLOCKS:
            self.current = None
        elif tag in _INLINE and self.styles:
            self.styles.pop()

    def handle_data(self, data):
        if self.current is None:
            if not data.strip():
                return
            self._block("p", [])
        self.current["runs"].append((data, set(filter(None, self.styles))))


def _parse(html: str):
    parser = _Blocks()
    parser.feed(html or "")
    parser.close()
    return parser.blocks


def html_text(html: str) -> str:
    lines = []
    for block in _parse(html):
        text = "".join(text for text, _ in block["runs"])
        lines.append(("• " if block["list"] == "ul" else "") + text)
    return "\n".join(lines).strip()


# ---------------------------------------------------------------- Word copy


def _full_width_picture(paragraph, image, margin):
    """An image across the whole page width, pulled out past the side margins."""
    paragraph.paragraph_format.left_indent = -margin
    paragraph.paragraph_format.right_indent = -margin
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(image), width=Cm(21.0))


def reference_lines(our_reference: str, their_reference: str | None) -> list:
    """The corner of the letter: our reference, and theirs when we are answering something."""
    lines = [("Our Ref: ", our_reference)]
    if their_reference:
        lines.append(("Your Ref: ", their_reference))
    return lines


def _letterhead_document(our_reference: str, their_reference: str | None = None):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    margin = Cm(2.1)
    section.left_margin = section.right_margin = margin
    # The header image is 4.8 cm tall at full width; text starts below it and the reference line.
    section.header_distance, section.top_margin = Cm(0), Cm(6.0)
    section.footer_distance, section.bottom_margin = Cm(0.2), Cm(2.4)
    style = document.styles["Normal"]
    # Every letter is set in Times New Roman.
    style.font.name = LETTER_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), LETTER_FONT)
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)

    header = section.header
    _full_width_picture(header.paragraphs[0], HEADER_PNG, margin)
    for index, (label_text, value_text) in enumerate(reference_lines(our_reference, their_reference)):
        reference = header.add_paragraph()
        reference.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        reference.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        label = reference.add_run(label_text)
        label.font.size = Pt(9)
        value = reference.add_run(value_text)
        value.bold = True
        value.font.size = Pt(9)
    _full_width_picture(section.footer.paragraphs[0], FOOTER_PNG, margin)
    return document


_ALIGN = {"ql-align-center": WD_ALIGN_PARAGRAPH.CENTER, "ql-align-right": WD_ALIGN_PARAGRAPH.RIGHT, "ql-align-justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
_SIZE = {"ql-size-small": Pt(10), "ql-size-large": Pt(15), "ql-size-huge": Pt(20)}
_HEADING_SIZE = {"h1": Pt(16), "h2": Pt(14), "h3": Pt(12)}


def html_to_docx(our_reference: str, html: str, footer_note: str | None = None, their_reference: str | None = None) -> bytes:
    """The letter as a Word document on letterhead."""
    document = _letterhead_document(our_reference, their_reference)
    for block in _parse(html):
        style = {"ul": "List Bullet", "ol": "List Number"}.get(block["list"])
        paragraph = document.add_paragraph(style=style) if style else document.add_paragraph()
        for cls in block["classes"]:
            if cls in _ALIGN:
                paragraph.alignment = _ALIGN[cls]
            elif cls.startswith("ql-indent-"):
                paragraph.paragraph_format.left_indent = Cm(1.27 * int(cls.rsplit("-", 1)[1]))
        for index, (text, marks) in enumerate(block["runs"]):
            if text == "\n" and index == len(block["runs"]) - 1:
                continue  # the editor marks an empty line with a trailing <br>; Word needs only the paragraph
            if text == "\n":
                (paragraph.runs[-1] if paragraph.runs else paragraph.add_run()).add_break(WD_BREAK.LINE)
                continue
            run = paragraph.add_run(text)
            run.bold = "b" in marks or block["tag"] in _HEADING_SIZE or None
            run.italic = "i" in marks or None
            run.underline = "u" in marks or None
            run.font.strike = "s" in marks or None
            size = next((_SIZE[m] for m in marks if m in _SIZE), _HEADING_SIZE.get(block["tag"]))
            if size:
                run.font.size = size
    if footer_note:
        document.add_paragraph()
        run = document.add_paragraph().add_run(footer_note)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
